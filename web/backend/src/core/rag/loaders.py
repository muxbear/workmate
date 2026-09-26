"""文档加载器——策略模式 + 组合模式。

每种文档类型对应一个策略类，FallbackLoaderStrategy 组合多个策略形成优先级链。
"""

import logging
import os
from abc import ABC, abstractmethod
from typing import Any, cast

from langchain_core.documents import Document

from core.rag.vision import OCR_PROMPT, VisionClient

logger = logging.getLogger(__name__)


def _clean_cell(text: str) -> str:
    """单元格文本归一：换行折成空格、竖线转义。

    竖线必须转义——否则单元格里的 ``|`` 会把 Markdown 表格的列切错，整张表读起来
    就乱了（而表格正是最需要保留结构的内容）。
    """
    return " ".join((text or "").split()).replace("|", "\\|")


def table_to_markdown(table: Any) -> str:
    """把 Word/PowerPoint 的表格渲染成 Markdown 表格。

    **为什么值得保留**：需求/接口文档里最要紧的往往就是那张对照表（"参数 → 默认值 →
    说明"）。纯文本抽取会把它压成一串没有标签的词，切片后既不可检索也不可读；
    Markdown 表格既让模型看得懂，也让人能读。

    python-docx 与 python-pptx 的 ``table.rows[].cells[].text`` 接口一致，因此
    Word 与 PowerPoint 共用这一份实现。
    """
    rows: list[list[str]] = [
        [_clean_cell(cell.text) for cell in row.cells] for row in table.rows
    ]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header = "| " + " | ".join(rows[0]) + " |"
    separator = "| " + " | ".join(["---"] * width) + " |"
    body = ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join([header, separator, *body])


class DocumentLoaderStrategy(ABC):
    """文档加载策略抽象接口。"""

    @abstractmethod
    def load(self, file_path: str) -> list[Document]:
        """加载文档并返回 LangChain Document 列表。"""
        ...


class OpenDataLoaderPDFStrategy(DocumentLoaderStrategy):
    """PDF 加载——langchain-opendataloader-pdf（优先策略）。"""

    def load(self, file_path: str) -> list[Document]:
        if not os.path.isfile(file_path):
            raise FileNotFoundError(f"PDF 文件未找到: {file_path}")
        from langchain_opendataloader_pdf import OpenDataLoaderPDFLoader
        docs = OpenDataLoaderPDFLoader(
                    file_path=str(file_path),
                    format="markdown",
                    table_method="cluster",
                    include_header_footer=True,
                    image_output="external",
                    image_dir="./images",
                    image_format="png"
                ).load()
        if not docs:
            raise RuntimeError(
                f"OpenDataLoaderPDFLoader 加载文件 {file_path} 返回空的 Documents。"
            )
        return docs


class PyPDFLoaderStrategy(DocumentLoaderStrategy):
    """PDF 加载——langchain_community PyPDFLoader（备选策略）。"""

    def load(self, file_path: str) -> list[Document]:
        from langchain_community.document_loaders import PyPDFLoader
        return PyPDFLoader(file_path).load()


class DocxLoaderStrategy(DocumentLoaderStrategy):
    """Word 文档加载——**保留表格**（迭代 6 T6.4）。

    此前用 ``Docx2txtLoader``：它只抽正文，**表格与图片都被静默丢掉**——用户传了
    一份带对照表的方案，入库后那张表根本不在了，而界面上不会有任何提示。

    现在用 python-docx 按**正文顺序**（``iter_inner_content``）遍历段落与表格，
    表格渲染成 Markdown。图片仍只计数（VLM 说明见后续批次）。
    """

    def load(self, file_path: str) -> list[Document]:
        from docx import Document as DocxDocument
        from docx.table import Table as DocxTable

        document = DocxDocument(file_path)
        parts: list[str] = []
        table_count = 0
        for item in document.iter_inner_content():
            if isinstance(item, DocxTable):
                rendered = table_to_markdown(item)
                if rendered:
                    parts.append(rendered)
                    table_count += 1
                continue
            text = (item.text or "").strip()
            if text:
                parts.append(text)

        image_count = len(document.inline_shapes)
        if image_count:
            # 不静默丢弃：正文里留一行明说"这里有几张图没进正文"，
            # 让"检索不到图里的内容"变成一个可解释的现象
            parts.append(f"（本文档含 {image_count} 张图片，其文字内容未提取）")
        if not parts:
            raise RuntimeError(f"No text content extracted from {file_path}")
        logger.debug(
            "DOCX 解析完成：%d 段/表（其中表格 %d），图片 %d",
            len(parts), table_count, image_count,
        )
        return [Document(
            page_content="\n\n".join(parts),
            metadata={
                "source": str(file_path),
                "tables": table_count,
                "images": image_count,
            },
        )]


class UnstructuredExcelStrategy(DocumentLoaderStrategy):
    """Excel 表格加载。"""

    def load(self, file_path: str) -> list[Document]:
        from langchain_community.document_loaders import UnstructuredExcelLoader
        return UnstructuredExcelLoader(file_path).load()


class PythonPPTXLoaderStrategy(DocumentLoaderStrategy):
    """PowerPoint 加载——使用 python-pptx（跨平台，无需 unstructured）。"""

    def load(self, file_path: str) -> list[Document]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs = Presentation(file_path)
        docs: list[Document] = []
        for i, slide in enumerate(prs.slides):
            texts: list[str] = []
            pictures = 0
            for shape in slide.shapes:
                if shape.has_text_frame:
                    tf = cast(Any, shape).text_frame
                    for para in tf.paragraphs:
                        t = para.text.strip()
                        if t:
                            texts.append(t)
                    continue
                # 表格此前被整块丢掉（只处理了 has_text_frame）——幻灯片里的表格
                # 往往就是全篇的结论
                if getattr(shape, "has_table", False):
                    rendered = table_to_markdown(cast(Any, shape).table)
                    if rendered:
                        texts.append(rendered)
                    continue
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    pictures += 1
            if pictures:
                texts.append(f"（本页含 {pictures} 张图片，其文字内容未提取）")
            if texts:
                docs.append(Document(
                    page_content="\n\n".join(texts),
                    metadata={"slide": i, "source": str(file_path), "images": pictures},
                ))
        if not docs:
            raise RuntimeError(f"No text content extracted from {file_path}")
        return docs


class UnstructuredPPTStrategy(DocumentLoaderStrategy):
    """PowerPoint 加载——unstructured（非 Windows 备选，Windows 上会崩溃）。"""

    def load(self, file_path: str) -> list[Document]:
        import platform
        if platform.system() == "Windows":
            raise RuntimeError(
                "UnstructuredPowerPointLoader is not supported on Windows "
                "due to python-magic incompatibility"
            )
        from langchain_community.document_loaders import UnstructuredPowerPointLoader
        return UnstructuredPowerPointLoader(file_path).load()


class CSVLoaderStrategy(DocumentLoaderStrategy):
    """CSV 加载。"""

    def load(self, file_path: str) -> list[Document]:
        from langchain_community.document_loaders import CSVLoader
        return CSVLoader(file_path).load()


class JSONLoaderStrategy(DocumentLoaderStrategy):
    """JSON 加载。"""

    def load(self, file_path: str) -> list[Document]:
        from langchain_community.document_loaders import JSONLoader
        return JSONLoader(file_path, jq_schema=".", text_content=False).load()


class MarkdownLoaderStrategy(DocumentLoaderStrategy):
    """Markdown 加载。"""

    def load(self, file_path: str) -> list[Document]:
        from langchain_community.document_loaders import UnstructuredMarkdownLoader
        return UnstructuredMarkdownLoader(file_path, mode="elements").load()


class HTMLLoaderStrategy(DocumentLoaderStrategy):
    """HTML 加载——使用 BeautifulSoup 解析并提取纯文本。"""

    def load(self, file_path: str) -> list[Document]:
        from langchain_community.document_loaders import BSHTMLLoader
        return BSHTMLLoader(file_path, open_encoding="utf-8").load()


class TextLoaderStrategy(DocumentLoaderStrategy):
    """纯文本加载。"""

    def load(self, file_path: str) -> list[Document]:
        from langchain_community.document_loaders import TextLoader
        return TextLoader(file_path, encoding="utf-8").load()


class ImageLoaderStrategy(DocumentLoaderStrategy):
    """图片加载——走「模型」页面配置的视觉模型做 OCR（迭代 6 T6.4）。

    此前依赖本机安装 Tesseract，未装即一律失败。而 ``png``/``jpg``/``jpeg`` 都在上传
    白名单里，于是**每一次图片上传都必然在解析阶段报错**，且错误信息指向一个用户在
    界面里无从解决的系统依赖。现在改为用模型页配置的 ``vision`` 模型：配了就能用，
    没配则给出可操作的提示。
    """

    def __init__(self, ocr: VisionClient | None = None) -> None:
        self._ocr = ocr

    def load(self, file_path: str) -> list[Document]:
        if self._ocr is None:
            raise RuntimeError(
                "图片解析需要视觉模型：请到「模型」页面添加 type=vision 的模型，"
                "并在本知识库的索引配置中开启 OCR 后重试。"
            )
        with open(file_path, "rb") as fh:
            data = fh.read()
        text = self._ocr.describe(OCR_PROMPT, data)
        if not text:
            # 不返回空文档：那会让文档以 「indexed、0 切片」的假成功入库
            raise RuntimeError("图片中未提取到文字（可能是纯图形、空白或过于模糊的图片）。")
        return [Document(page_content=text, metadata={"source": str(file_path)})]


class FallbackLoaderStrategy(DocumentLoaderStrategy):
    """组合策略——按优先级链依次尝试，返回第一个成功的结果。"""

    def __init__(self, strategies: list[DocumentLoaderStrategy]):
        self._strategies = strategies

    def load(self, file_path: str) -> list[Document]:
        errors: list[str] = []
        for strategy in self._strategies:
            try:
                return strategy.load(file_path)
            except Exception as e:
                errors.append(f"{type(strategy).__name__}: {e}")
                logger.debug("Loader %s failed: %s", type(strategy).__name__, e)
        raise ValueError(f"所有加载器出现异常: {'; '.join(errors)}")


class DocumentLoaderRegistry:
    """文档加载策略注册表。"""

    def __init__(self):
        self._strategies: dict[str, DocumentLoaderStrategy] = {}

    def register(self, file_type: str, strategy: DocumentLoaderStrategy) -> None:
        self._strategies[file_type] = strategy

    def get_strategy(self, file_type: str) -> DocumentLoaderStrategy:
        if file_type not in self._strategies:
            raise ValueError(f"Unsupported file type: {file_type}")
        return self._strategies[file_type]

    def load(self, file_path: str, file_type: str) -> list[Document]:
        return self.get_strategy(file_type).load(file_path)


def create_default_loader_registry(
    ocr: VisionClient | None = None,
) -> DocumentLoaderRegistry:
    """创建预注册所有内置文件类型的加载器注册表。

    Args:
        ocr: 视觉模型客户端。为 ``None`` 时图片类文档的解析会明确失败并提示去配置
            模型——**不会**静默产出一个空文档。未开启 OCR 的知识库继续沿用启动时
            构建的那一份（见 ``api.knowledge_base.facade``），零开销。
    """
    registry = DocumentLoaderRegistry()

    # PDF: 优先 opendataloader_pdf，备选 PyPDFLoader
    registry.register("pdf", FallbackLoaderStrategy([
        OpenDataLoaderPDFStrategy(),
        PyPDFLoaderStrategy(),
    ]))

    registry.register("docx", DocxLoaderStrategy())
    registry.register("xlsx", UnstructuredExcelStrategy())
    registry.register("pptx", FallbackLoaderStrategy([
        PythonPPTXLoaderStrategy(),
        UnstructuredPPTStrategy(),
    ]))
    registry.register("csv", CSVLoaderStrategy())
    registry.register("json", JSONLoaderStrategy())
    registry.register("md", MarkdownLoaderStrategy())
    registry.register("html", HTMLLoaderStrategy())
    registry.register("txt", TextLoaderStrategy())

    image_strategy = ImageLoaderStrategy(ocr)
    for ft in ("png", "jpg", "jpeg", "image"):
        registry.register(ft, image_strategy)

    return registry
