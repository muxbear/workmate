"""插图说明：DOCX / PPTX / PDF 里的图转成可检索的文字（迭代 6 T6.4 批次④）。

此前图只被计数，正文里留一句「本文档含 N 张图片，其文字内容未提取」——那是批次①
刻意留下的占位（"让检索不到图里的内容变成一个可解释的现象"）。本批次把它补上。

要紧的不是"能不能生成说明"，而是**说明落在正文的哪个位置**：统一追加到文末会让
"图 1 的说明"排到第 40 页后面，检索到它的人无从判断它属于哪里。顺序是硬约束，
所以用例大量断言相对位置而不是只断言存在。
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import httpx
import pytest
from docx import Document as DocxDocument
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu

from core.rag.loaders import (
    DocxLoaderStrategy,
    OpenDataLoaderPDFStrategy,
    PythonPPTXLoaderStrategy,
)
from core.rag.ocr import LOAD_REPORT_KEY, build_load_report, format_ocr_warning
from core.rag.vision import CAPTION_PROMPT, ImageCaptioner, VisionClient


def make_png(color: tuple[int, int, int] = (200, 30, 30)) -> bytes:
    image = Image.new("RGB", (80, 60), "white")
    ImageDraw.Draw(image).rectangle([8, 8, 72, 52], fill=color)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


class RecordingTransport(httpx.BaseTransport):
    """按脚本回说明，并记下每次请求用的提示词。"""

    def __init__(self, answers: list[str] | None = None) -> None:
        self.answers = list(answers or [])
        self.calls = 0
        self.prompts: list[str] = []

    def handle_request(self, request: httpx.Request) -> httpx.Response:
        self.calls += 1
        body = json.loads(request.content.decode("utf-8"))
        self.prompts.append(body["messages"][0]["content"][1]["text"])
        answer = self.answers.pop(0) if self.answers else f"第 {self.calls} 张图的说明"
        return httpx.Response(200, json={"choices": [{"message": {"content": answer}}]})


class FailingTransport(httpx.BaseTransport):
    def handle_request(self, request: httpx.Request) -> httpx.Response:
        return httpx.Response(500, json={"error": "boom"})


def make_ocr(transport: httpx.BaseTransport) -> VisionClient:
    return VisionClient(
        model="qwen-vl-ocr",
        api_base="https://dashscope.aliyuncs.com/compatible-mode/v1",
        api_key="sk-test",
        transport=transport,
    )


def make_captioner(
    transport: httpx.BaseTransport,
    **kwargs,
) -> ImageCaptioner:
    return ImageCaptioner(make_ocr(transport), **kwargs)


def make_docx(path: Path, *, images: int = 2) -> Path:
    """段落 / 图 / 表格 / 段落 / 图 —— 用来验证说明落在正确的相对位置。"""
    doc = DocxDocument()
    doc.add_paragraph("第一段文字")
    for i in range(images):
        doc.add_picture(io.BytesIO(make_png((200 - i * 60, 30, 30))))
        if i == 0:
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "表格左"
            table.cell(0, 1).text = "表格右"
            doc.add_paragraph("中间段落")
    doc.add_paragraph("最后一段文字")
    doc.save(str(path))
    return path


def make_pptx(path: Path, *, slides: int = 2) -> Path:
    prs = Presentation()
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"第 {i + 1} 页标题"
        slide.shapes.add_picture(
            io.BytesIO(make_png((30, 120 - i * 40, 200))),
            Emu(1000000),
            Emu(2000000),
        )
    prs.save(str(path))
    return path


# ─── DOCX ───────────────────────────────────────────────────────────────────


class TestDocxCaptions:
    def test_captions_are_inserted_in_body_order(self, tmp_path: Path):
        """顺序是硬约束：说明必须落在**图所在的位置**，不能统一追加到文末。

        `iter_inner_content` 遍历段落与表格就是为了保正文顺序；说明若堆到文末，
        检索到"图 1 的说明"的人无从判断它属于哪一段。
        """
        target = make_docx(tmp_path / "带图.docx")
        transport = RecordingTransport(["红色示意图", "蓝色示意图"])
        docs = DocxLoaderStrategy(make_captioner(transport)).load(str(target))
        text = docs[0].page_content

        assert transport.calls == 2
        assert text.index("第一段文字") < text.index("红色示意图")
        assert text.index("红色示意图") < text.index("| 表格左 |")
        assert text.index("| 表格左 |") < text.index("中间段落")
        assert text.index("中间段落") < text.index("蓝色示意图")
        assert text.index("蓝色示意图") < text.index("最后一段文字")

    def test_caption_prompt_is_the_caption_one_not_the_ocr_one(self, tmp_path: Path):
        """说明用「描述这张图」，OCR 用「只输出图中文字」——两者目标不同，别弄反。"""
        target = make_docx(tmp_path / "带图.docx", images=1)
        transport = RecordingTransport()
        DocxLoaderStrategy(make_captioner(transport)).load(str(target))
        assert transport.prompts == [CAPTION_PROMPT]

    def test_without_captioner_behaviour_is_unchanged(self, tmp_path: Path):
        """没配视觉模型时退回批次①的行为：计数占位，不静默丢弃。"""
        target = make_docx(tmp_path / "带图.docx")
        text = DocxLoaderStrategy().load(str(target))[0].page_content

        assert "含 2 张图片" in text
        assert "未提取" in text

    def test_placeholder_only_counts_the_images_that_failed(self, tmp_path: Path):
        """拿到说明的图不该再留"未提取"占位——既解释了又说不清，自相矛盾。"""
        target = make_docx(tmp_path / "带图.docx")
        text = (
            DocxLoaderStrategy(
                make_captioner(FailingTransport()),
            )
            .load(str(target))[0]
            .page_content
        )

        assert "含 2 张图片" in text

    def test_failed_caption_falls_back_to_the_placeholder(self, tmp_path: Path):
        target = make_docx(tmp_path / "带图.docx")
        docs = DocxLoaderStrategy(make_captioner(FailingTransport())).load(str(target))

        report = docs[0].metadata[LOAD_REPORT_KEY]
        assert report["images"] == 2
        assert report["images_captioned"] == 0
        assert report["images_failed"] == 2

    def test_report_counts_captioned_images(self, tmp_path: Path):
        target = make_docx(tmp_path / "带图.docx")
        docs = DocxLoaderStrategy(
            make_captioner(RecordingTransport()),
        ).load(str(target))

        report = docs[0].metadata[LOAD_REPORT_KEY]
        assert report["images"] == 2
        assert report["images_captioned"] == 2
        assert "未提取" not in docs[0].page_content

    def test_budget_caps_the_number_of_calls(self, tmp_path: Path):
        """一份配图很多的文档不能无限烧调用——超预算的图记数跳过，不静默丢。"""
        target = make_docx(tmp_path / "带图.docx")
        transport = RecordingTransport()
        docs = DocxLoaderStrategy(
            make_captioner(transport, max_images=1),
        ).load(str(target))

        assert transport.calls == 1
        report = docs[0].metadata[LOAD_REPORT_KEY]
        assert report["images_captioned"] == 1
        assert report["images_skipped"] == 1
        assert "含 1 张图片" in docs[0].page_content


# ─── PPTX ───────────────────────────────────────────────────────────────────


class TestPptxCaptions:
    def test_each_slide_gets_its_own_caption(self, tmp_path: Path):
        target = make_pptx(tmp_path / "带图.pptx")
        transport = RecordingTransport(["第一页的图说明", "第二页的图说明"])
        docs = PythonPPTXLoaderStrategy(make_captioner(transport)).load(str(target))

        assert len(docs) == 2
        assert "第 1 页标题" in docs[0].page_content
        assert "第一页的图说明" in docs[0].page_content
        # 说明不能串页
        assert "第二页的图说明" not in docs[0].page_content
        assert "第二页的图说明" in docs[1].page_content

    def test_without_captioner_keeps_the_placeholder(self, tmp_path: Path):
        target = make_pptx(tmp_path / "带图.pptx", slides=1)
        docs = PythonPPTXLoaderStrategy().load(str(target))
        assert "含 1 张图片" in docs[0].page_content


# ─── 告警文案 ───────────────────────────────────────────────────────────────


class TestCaptionWarning:
    def test_all_captioned_says_nothing(self):
        assert (
            format_ocr_warning(
                build_load_report(images=5, images_captioned=5),
            )
            is None
        )

    def test_uncaptioned_images_are_reported(self):
        warning = format_ocr_warning(
            build_load_report(images=10, images_captioned=8, images_skipped=2),
        )
        assert warning is not None
        assert "2" in warning and "10" in warning

    def test_pages_and_images_are_reported_together(self):
        """两类缺失含义不同，要能同时看到。"""
        warning = format_ocr_warning(
            build_load_report(
                ocr_pages=3,
                ocr_skipped_budget=1,
                total_pages=4,
                images=2,
                images_failed=1,
            ),
        )
        assert warning is not None
        assert "页" in warning and "图" in warning


# ─── PDF（需要 Java，跑不了就跳过） ─────────────────────────────────────────


requires_java = pytest.mark.skipif(
    __import__("shutil").which("java") is None,
    reason="opendataloader 依赖 Java CLI，本机没有就跳过（备选路径 PyPDFLoader 不导出图片）",
)


def _make_text_pdf_with_figure(path: Path) -> Path:
    """一页里既有文本层又有插图——用来验证图片导出的落点与就地替换。"""
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    picture = path.parent / "_figure_only.pdf"
    image = Image.new("RGB", (400, 300), "white")
    ImageDraw.Draw(image).rectangle([20, 20, 380, 120], fill=(200, 30, 30))
    with open(picture, "wb") as fh:
        image.save(fh, "PDF", resolution=150)

    text_only = path.parent / "_text_only.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font_ref = writer._add_object(
        DictionaryObject(
            {  # noqa: SLF001
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref}),
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 18 Tf 60 760 Td (A page with a figure below) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)  # noqa: SLF001
    with open(text_only, "wb") as fh:
        writer.write(fh)

    merged = PdfWriter()
    target_page = merged.add_blank_page(width=595, height=842)
    target_page.merge_page(PdfReader(str(text_only)).pages[0])
    target_page.merge_page(PdfReader(str(picture)).pages[0])
    with open(path, "wb") as fh:
        merged.write(fh)
    return path


@requires_java
class TestPdfCaptions:
    def test_images_land_in_the_document_own_directory(self, tmp_path: Path):
        """导出目录必须是**绝对的、按文档隔离的**。

        此前写的是相对路径 ``./images``，按进程工作目录解析——所有文档共用同一个
        目录，而导出文件名是顺序计数（imageFile1.png…），并发的两份文档会互相覆盖
        对方的插图，最后这份的说明配的是别人的图。
        """
        target = _make_text_pdf_with_figure(tmp_path / "带图.pdf")
        OpenDataLoaderPDFStrategy(make_captioner(RecordingTransport())).load(
            str(target)
        )

        exported = tmp_path / "images"
        assert exported.is_dir(), "插图应导出到文档自己目录下的 images/"
        assert list(exported.glob("*.png")), "导出目录里应有图片"

    def test_caption_replaces_the_image_reference_in_place(self, tmp_path: Path):
        """引用标出了图在正文中的位置，就地替换天然保序。"""
        target = _make_text_pdf_with_figure(tmp_path / "带图.pdf")
        docs = OpenDataLoaderPDFStrategy(
            make_captioner(RecordingTransport(["一张红色示意图"])),
        ).load(str(target))

        text = "\n".join(d.page_content for d in docs)
        assert "一张红色示意图" in text
        assert "![" not in text, "图片引用应已被说明替换"

    def test_failed_caption_keeps_the_reference(self, tmp_path: Path):
        """拿不到说明时保留原引用——它至少标出了"这里有张图"，抹掉才是真丢失。"""
        target = _make_text_pdf_with_figure(tmp_path / "带图.pdf")
        docs = OpenDataLoaderPDFStrategy(
            make_captioner(FailingTransport()),
        ).load(str(target))

        text = "\n".join(d.page_content for d in docs)
        assert "![" in text
        assert docs[0].metadata[LOAD_REPORT_KEY]["images_failed"] >= 1

    def test_without_captioner_references_are_untouched(self, tmp_path: Path):
        target = _make_text_pdf_with_figure(tmp_path / "带图.pdf")
        docs = OpenDataLoaderPDFStrategy().load(str(target))
        assert "![" in "\n".join(d.page_content for d in docs)

    def test_caption_counts_survive_the_page_ocr_layer(self, tmp_path: Path):
        """回归：外层"扫描页补 OCR"的报告不能把内层的图片计数抹掉。

        数字 PDF 走的正是外层"没有扫描页"那条早返回路径——那里若直接赋值覆盖，
        "有几张图没生成说明"就永远不会出现在界面上，而插图恰恰是数字 PDF 才有的。

        用例必须给一个**非空**的 ocr 客户端：为 None 时外层在更早的地方就返回了，
        根本走不到挂报告那一步，用例会空转通过。
        """
        from core.rag.loaders import FallbackLoaderStrategy, OcrAugmentedPdfStrategy

        target = _make_text_pdf_with_figure(tmp_path / "带图.pdf")
        page_transport = RecordingTransport()
        strategy = OcrAugmentedPdfStrategy(
            FallbackLoaderStrategy(
                [
                    OpenDataLoaderPDFStrategy(make_captioner(FailingTransport())),
                ]
            ),
            ocr=make_ocr(page_transport),
        )
        docs = strategy.load(str(target))

        # 数字 PDF 不该产生任何页面识别调用（判据严到"至少得有一个词"）
        assert page_transport.calls == 0
        report = docs[0].metadata[LOAD_REPORT_KEY]
        assert report["images_failed"] >= 1
        warning = format_ocr_warning(report)
        assert warning is not None and "插图" in warning
